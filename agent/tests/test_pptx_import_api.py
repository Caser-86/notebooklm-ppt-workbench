from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.main import app
from app.services.artifacts import ARTIFACTS_ROOT


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def test_upload_pptx_creates_import_record(build_fixture_pptx):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "PPTX Import", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx("demo.pptx", ["Editable rebuild"])

    with pptx_path.open("rb") as handle:
        response = client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("demo.pptx", handle, PPTX_MIME)},
        )

    assert response.status_code == 201
    body = response.json()
    assert body["project_id"] == project["id"]
    assert body["filename"] == "demo.pptx"
    assert body["status"] == "ready"
    assert body["page_count"] == 1


def test_imported_pptx_can_be_rebuilt_into_artifacts(build_fixture_pptx):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Import Flow", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx("flow.pptx", ["Editable rebuild"])

    with pptx_path.open("rb") as handle:
        imported = client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("flow.pptx", handle, PPTX_MIME)},
        ).json()

    response = client.post(f"/imports/{imported['id']}/rebuild")

    assert response.status_code == 200
    assert {artifact["id"] for artifact in response.json()["artifacts"]} == {"display-clone", "editable-rebuild"}


def test_imported_multi_slide_pptx_preserves_slide_count_in_rebuild(build_fixture_pptx):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Import Flow", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx("multi-slide.pptx", ["Editable rebuild", "NotebookLM export"])

    with pptx_path.open("rb") as handle:
        imported = client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("multi-slide.pptx", handle, PPTX_MIME)},
        ).json()

    rebuild = client.post(f"/imports/{imported['id']}/rebuild")

    assert rebuild.status_code == 200

    display_path = ARTIFACTS_ROOT / str(project["id"]) / "rebuild-001" / "display-clone.pptx"
    editable_path = ARTIFACTS_ROOT / str(project["id"]) / "rebuild-001" / "editable-rebuild.pptx"

    assert len(Presentation(display_path).slides) == 2
    assert len(Presentation(editable_path).slides) == 2


def test_imported_table_rebuild_stays_editable_table(build_fixture_pptx_with_table):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Table Import", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx_with_table("table-import.pptx")

    with pptx_path.open("rb") as handle:
        imported = client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("table-import.pptx", handle, PPTX_MIME)},
        ).json()

    rebuild = client.post(f"/imports/{imported['id']}/rebuild")

    assert rebuild.status_code == 200
    editable_path = ARTIFACTS_ROOT / str(project["id"]) / "rebuild-001" / "editable-rebuild.pptx"
    presentation = Presentation(editable_path)

    assert any(shape.has_table for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE)


def test_imported_chart_counts_appear_in_import_summary(build_fixture_pptx_with_chart):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Chart Import", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx_with_chart("chart-summary.pptx")

    with pptx_path.open("rb") as handle:
        client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("chart-summary.pptx", handle, PPTX_MIME)},
        )

    imports = client.get(f"/projects/{project['id']}/imports").json()

    assert imports[0]["object_summary"]["imported_chart"] == 1
    assert imports[0]["slide_assets"][0]["object_summary"]["imported_chart"] == 1


def test_imported_chart_rebuild_preserves_chart_presence(build_fixture_pptx_with_chart):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Chart Rebuild", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx_with_chart("chart-rebuild.pptx")

    with pptx_path.open("rb") as handle:
        imported = client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("chart-rebuild.pptx", handle, PPTX_MIME)},
        ).json()

    response = client.post(f"/imports/{imported['id']}/rebuild")

    assert response.status_code == 200
    editable = Presentation(ARTIFACTS_ROOT / str(project["id"]) / "rebuild-001" / "editable-rebuild.pptx")
    assert any(shape.has_table for shape in editable.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE)


def test_imported_group_counts_appear_in_import_summary(build_fixture_pptx_with_group_shape):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Group Import", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx_with_group_shape("group-summary.pptx")

    with pptx_path.open("rb") as handle:
        client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("group-summary.pptx", handle, PPTX_MIME)},
        )

    imports = client.get(f"/projects/{project['id']}/imports").json()

    assert imports[0]["object_summary"]["unsupported_group"] >= 1
    assert imports[0]["slide_assets"][0]["object_summary"]["unsupported_group"] >= 1


def test_unsupported_group_rebuild_preserves_group_presence(build_fixture_pptx_with_group_shape):
    client = TestClient(app)
    project = client.post("/projects", json={"title": "Group Rebuild", "preferred_language": "zh-CN"}).json()
    pptx_path = build_fixture_pptx_with_group_shape("group-rebuild.pptx")

    with pptx_path.open("rb") as handle:
        imported = client.post(
            f"/projects/{project['id']}/imports/pptx",
            files={"file": ("group-rebuild.pptx", handle, PPTX_MIME)},
        ).json()

    response = client.post(f"/imports/{imported['id']}/rebuild")

    assert response.status_code == 200
    editable = Presentation(ARTIFACTS_ROOT / str(project["id"]) / "rebuild-001" / "editable-rebuild.pptx")
    slide_texts = [shape.text for shape in editable.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]

    assert any("Unsupported grouped content" in text for text in slide_texts)
