import json
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from app.services.reconstruct.editable_rebuild import build_editable_rebuild


def test_build_editable_rebuild_creates_text_boxes_from_ocr_blocks(tmp_path):
    blocks = json.loads(Path("tests/fixtures/ocr_blocks.json").read_text(encoding="utf-8"))
    output_path = tmp_path / "editable-rebuild.pptx"

    build_editable_rebuild(blocks, output_path)

    assert output_path.exists()
    assert output_path.stat().st_size > 0


def test_build_editable_rebuild_creates_multiple_slides_and_groups_same_line_blocks(tmp_path):
    blocks = [
        {"text": "Quarterly", "slide_index": 0, "x": 1, "y": 1, "width": 1.6, "height": 0.5, "font_size": 24},
        {"text": "update", "slide_index": 0, "x": 2.7, "y": 1.02, "width": 1.6, "height": 0.5, "font_size": 24},
        {"text": "Revenue", "slide_index": 1, "x": 1, "y": 1, "width": 1.6, "height": 0.5, "font_size": 22},
        {"text": "22%", "slide_index": 1, "x": 2.8, "y": 1.01, "width": 1.2, "height": 0.5, "font_size": 22},
    ]
    output_path = tmp_path / "editable-rebuild-multi.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    slide_texts = [
        [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        for slide in presentation.slides
    ]

    assert len(presentation.slides) == 2
    assert slide_texts[0] == ["Quarterly update"]
    assert slide_texts[1] == ["Revenue 22%"]


def test_build_editable_rebuild_marks_title_blocks_as_bolder_than_body(tmp_path):
    blocks = [
        {"text": "Market expansion", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.5, "height": 0.7, "font_size": 30},
        {"text": "We entered three new regions this quarter.", "slide_index": 0, "x": 1, "y": 2.1, "width": 6.2, "height": 0.6, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-title-body.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]

    title_run = shapes_with_text[0].text_frame.paragraphs[0].runs[0]
    body_run = shapes_with_text[1].text_frame.paragraphs[0].runs[0]

    assert title_run.text == "Market expansion"
    assert title_run.font.bold is True
    assert body_run.font.bold in (None, False)
    assert title_run.font.size.pt > body_run.font.size.pt


def test_build_editable_rebuild_turns_bullet_lines_into_indented_list_items(tmp_path):
    blocks = [
        {"text": "Plan", "slide_index": 0, "x": 1, "y": 0.8, "width": 3.4, "height": 0.6, "font_size": 28},
        {"text": "- First checkpoint", "slide_index": 0, "x": 1.2, "y": 2.0, "width": 4.2, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-list.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]

    list_paragraph = shapes_with_text[1].text_frame.paragraphs[0]

    assert list_paragraph.level == 1
    assert list_paragraph.runs[0].text == "First checkpoint"


def test_build_editable_rebuild_preserves_space_below_titles(tmp_path):
    blocks = [
        {"text": "Growth plan", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 30},
        {"text": "Body text begins too close in the OCR output.", "slide_index": 0, "x": 1, "y": 1.0, "width": 6, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-spacing.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]
    title_shape = shapes_with_text[0]
    body_shape = shapes_with_text[1]

    assert body_shape.top >= title_shape.top + title_shape.height + Inches(0.18)


def test_build_editable_rebuild_keeps_consecutive_list_items_on_same_indent(tmp_path):
    blocks = [
        {"text": "Checklist", "slide_index": 0, "x": 1, "y": 0.8, "width": 3.4, "height": 0.6, "font_size": 28},
        {"text": "- First item", "slide_index": 0, "x": 1.2, "y": 2.0, "width": 4.2, "height": 0.45, "font_size": 18},
        {"text": "- Second item", "slide_index": 0, "x": 1.48, "y": 2.46, "width": 4.0, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-list-indent.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]
    list_shape = shapes_with_text[1]
    list_paragraphs = [paragraph for paragraph in list_shape.text_frame.paragraphs if paragraph.text.strip()]

    assert len(shapes_with_text) == 2
    assert [paragraph.level for paragraph in list_paragraphs] == [1, 1]


def test_build_editable_rebuild_caps_body_width_for_readability(tmp_path):
    blocks = [
        {"text": "Body heading", "slide_index": 0, "x": 1, "y": 0.8, "width": 3.5, "height": 0.6, "font_size": 28},
        {
            "text": "A very wide paragraph should be narrowed to a more readable text column in the rebuilt slide.",
            "slide_index": 0,
            "x": 1,
            "y": 2.0,
            "width": 10.5,
            "height": 0.6,
            "font_size": 18,
        },
    ]
    output_path = tmp_path / "editable-rebuild-body-width.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]
    body_shape = shapes_with_text[1]

    assert body_shape.width <= Inches(8.2)


def test_build_editable_rebuild_merges_consecutive_body_blocks_into_one_text_frame(tmp_path):
    blocks = [
        {"text": "Expansion update", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 30},
        {"text": "We launched three pilots in Southeast Asia.", "slide_index": 0, "x": 1, "y": 2.0, "width": 6.1, "height": 0.45, "font_size": 18},
        {"text": "Each pilot now has local distribution support.", "slide_index": 0, "x": 1.02, "y": 2.55, "width": 6.0, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-body-paragraphs.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]
    body_shape = shapes_with_text[1]
    body_paragraphs = [paragraph.text for paragraph in body_shape.text_frame.paragraphs if paragraph.text.strip()]

    assert len(shapes_with_text) == 2
    assert body_paragraphs == [
        "We launched three pilots in Southeast Asia.",
        "Each pilot now has local distribution support.",
    ]


def test_build_editable_rebuild_merges_consecutive_list_items_into_one_text_frame(tmp_path):
    blocks = [
        {"text": "Execution plan", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 30},
        {"text": "- Confirm export", "slide_index": 0, "x": 1.2, "y": 2.0, "width": 4.2, "height": 0.45, "font_size": 18},
        {"text": "- Trigger rebuild", "slide_index": 0, "x": 1.44, "y": 2.5, "width": 4.0, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-list-paragraphs.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    shapes_with_text = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]
    list_shape = shapes_with_text[1]
    list_paragraphs = [paragraph for paragraph in list_shape.text_frame.paragraphs if paragraph.text.strip()]

    assert len(shapes_with_text) == 2
    assert [paragraph.text for paragraph in list_paragraphs] == ["Confirm export", "Trigger rebuild"]
    assert [paragraph.level for paragraph in list_paragraphs] == [1, 1]


def test_build_editable_rebuild_preserves_image_blocks_and_keeps_body_outside_image(tmp_path):
    blocks = [
        {"text": "Launch overview", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 30},
        {
            "text": "",
            "slide_index": 0,
            "x": 7.2,
            "y": 1.9,
            "width": 4.0,
            "height": 2.6,
            "content_type": "image",
            "image_path": "tests/fixtures/slides/slide-1.png",
        },
        {
            "text": "Supporting narrative should stay clear of the image region on the right side of the slide.",
            "slide_index": 0,
            "x": 1,
            "y": 2.0,
            "width": 10.2,
            "height": 0.6,
            "font_size": 18,
        },
    ]
    output_path = tmp_path / "editable-rebuild-image-layout.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    slide = presentation.slides[0]
    picture_shapes = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    text_shapes = [shape for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]

    assert len(picture_shapes) == 1
    body_shape = text_shapes[1]
    picture_shape = picture_shapes[0]

    assert body_shape.left + body_shape.width <= picture_shape.left - Inches(0.08)


def test_build_editable_rebuild_formats_captions_as_secondary_text(tmp_path):
    blocks = [
        {"text": "Growth snapshot", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 30},
        {
            "text": "",
            "slide_index": 0,
            "x": 6.8,
            "y": 1.5,
            "width": 4.0,
            "height": 2.2,
            "content_type": "image",
            "image_path": "tests/fixtures/slides/slide-1.png",
        },
        {
            "text": "Pilot launch event, March 2026",
            "slide_index": 0,
            "x": 6.9,
            "y": 3.82,
            "width": 3.5,
            "height": 0.35,
            "font_size": 12,
        },
    ]
    output_path = tmp_path / "editable-rebuild-caption-style.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    slide = presentation.slides[0]
    picture_shape = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    text_shapes = [shape for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
    caption_shape = text_shapes[-1]
    caption_run = caption_shape.text_frame.paragraphs[0].runs[0]

    assert caption_shape.left == picture_shape.left
    assert caption_run.font.italic is True
    assert caption_run.font.bold in (None, False)


def test_build_editable_rebuild_uses_two_column_width_limits_for_body_blocks(tmp_path):
    blocks = [
        {"text": "Execution review", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 30},
        {"text": "Left column summary", "slide_index": 0, "x": 1, "y": 2.0, "width": 6.0, "height": 0.45, "font_size": 18},
        {"text": "Right column summary", "slide_index": 0, "x": 7.1, "y": 2.02, "width": 5.8, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-two-column-widths.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    text_shapes = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and shape.text.strip()]
    left_shape = text_shapes[1]
    right_shape = text_shapes[2]

    assert left_shape.width <= Inches(5.1)
    assert right_shape.left >= Inches(6.8)
    assert right_shape.width <= Inches(5.1)


def test_build_editable_rebuild_renders_detected_tables_as_ppt_tables(tmp_path):
    blocks = [
        {"text": "Quarterly metrics", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.4, "height": 0.7, "font_size": 28},
        {"text": "Region", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 4.0, "y": 2.0, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 4.0, "y": 2.7, "width": 2.1, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    slide = presentation.slides[0]
    table_shapes = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE]

    assert len(table_shapes) == 1
    table = table_shapes[0].table
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Region"
    assert table.cell(0, 1).text == "Revenue"
    assert table.cell(1, 0).text == "APAC"
    assert table.cell(1, 1).text == "$2.4M"


def test_build_editable_rebuild_styles_table_header_row(tmp_path):
    blocks = [
        {"text": "Quarterly metrics", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.4, "height": 0.7, "font_size": 28},
        {"text": "Region", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 4.0, "y": 2.0, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 4.0, "y": 2.7, "width": 2.1, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-header-style.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    header_run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    body_run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]

    assert header_run.font.bold is True
    assert body_run.font.bold in (None, False)


def test_build_editable_rebuild_applies_table_header_fill_and_grid_borders(tmp_path):
    blocks = [
        {"text": "Quarterly metrics", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.4, "height": 0.7, "font_size": 28},
        {"text": "Region", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 4.0, "y": 2.0, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.1, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 4.0, "y": 2.7, "width": 2.1, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-style.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    header_fill = table.cell(0, 0).fill.fore_color.rgb
    body_fill = table.cell(1, 0).fill.fore_color.rgb
    cell_xml = table.cell(0, 0)._tc.xml

    assert header_fill == RGBColor(0xE9, 0xEE, 0xF7)
    assert body_fill != header_fill
    assert "a:lnL" in cell_xml
    assert "a:lnR" in cell_xml
    assert "a:lnT" in cell_xml
    assert "a:lnB" in cell_xml


def test_build_editable_rebuild_sets_table_column_widths_from_detected_grid(tmp_path):
    blocks = [
        {"text": "Pipeline status", "slide_index": 0, "x": 1, "y": 0.8, "width": 4.2, "height": 0.7, "font_size": 28},
        {"text": "Stage", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 3.1, "height": 0.45, "font_size": 18},
        {"text": "Count", "slide_index": 0, "x": 4.7, "y": 2.0, "width": 1.2, "height": 0.45, "font_size": 18},
        {"text": "Qualified", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 3.1, "height": 0.45, "font_size": 18},
        {"text": "18", "slide_index": 0, "x": 4.7, "y": 2.7, "width": 1.2, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-widths.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert table.columns[0].width > table.columns[1].width


def test_build_editable_rebuild_merges_header_cells_across_columns(tmp_path):
    blocks = [
        {"text": "Regional performance", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.2, "height": 0.6, "font_size": 26},
        {"text": "Q1 2026", "slide_index": 0, "x": 1.0, "y": 1.5, "width": 4.9, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 3.1, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 4.7, "y": 2.0, "width": 1.2, "height": 0.45, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 3.1, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 4.7, "y": 2.7, "width": 1.2, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-colspan.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Q1 2026"
    assert table.cell(1, 0).text == "Region"
    assert table.cell(1, 1).text == "Revenue"
    assert table.cell(2, 0).text == "APAC"
    assert table.cell(2, 1).text == "$2.4M"
    assert 'gridSpan="2"' in table.cell(0, 0)._tc.xml


def test_build_editable_rebuild_merges_first_column_cells_across_rows(tmp_path):
    blocks = [
        {"text": "Segment performance", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.2, "height": 0.6, "font_size": 26},
        {"text": "Category", "slide_index": 0, "x": 1.0, "y": 1.6, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 3.5, "y": 1.6, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 5.7, "y": 1.6, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.3, "width": 2.0, "height": 1.0, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 3.5, "y": 2.3, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 5.7, "y": 2.3, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "EMEA", "slide_index": 0, "x": 3.5, "y": 2.9, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "$1.8M", "slide_index": 0, "x": 5.7, "y": 2.9, "width": 1.4, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-rowspan.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert table.cell(1, 0).text == "Enterprise"
    assert table.cell(1, 1).text == "APAC"
    assert table.cell(1, 2).text == "$2.4M"
    assert table.cell(2, 1).text == "EMEA"
    assert table.cell(2, 2).text == "$1.8M"
    assert 'rowSpan="2"' in table.cell(1, 0)._tc.xml


def test_build_editable_rebuild_merges_first_column_cells_across_three_rows(tmp_path):
    blocks = [
        {"text": "Segment performance", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.2, "height": 0.6, "font_size": 26},
        {"text": "Category", "slide_index": 0, "x": 1.0, "y": 1.6, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 3.5, "y": 1.6, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 5.7, "y": 1.6, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.3, "width": 2.0, "height": 1.65, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 3.5, "y": 2.3, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 5.7, "y": 2.3, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "EMEA", "slide_index": 0, "x": 3.5, "y": 2.9, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "$1.8M", "slide_index": 0, "x": 5.7, "y": 2.9, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "Americas", "slide_index": 0, "x": 3.5, "y": 3.5, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "$3.1M", "slide_index": 0, "x": 5.7, "y": 3.5, "width": 1.4, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-rowspan-three.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 4
    assert len(table.columns) == 3
    assert table.cell(1, 0).text == "Enterprise"
    assert table.cell(1, 1).text == "APAC"
    assert table.cell(2, 1).text == "EMEA"
    assert table.cell(3, 1).text == "Americas"
    assert table.cell(3, 2).text == "$3.1M"
    assert 'rowSpan="3"' in table.cell(1, 0)._tc.xml


def test_build_editable_rebuild_merges_non_first_column_cells_across_rows(tmp_path):
    blocks = [
        {"text": "Operations review", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.2, "height": 0.6, "font_size": 26},
        {"text": "Category", "slide_index": 0, "x": 1.0, "y": 1.6, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 3.5, "y": 1.6, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 5.7, "y": 1.6, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.3, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "APAC + EMEA", "slide_index": 0, "x": 3.5, "y": 2.3, "width": 1.7, "height": 1.0, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 5.7, "y": 2.3, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.9, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "$1.8M", "slide_index": 0, "x": 5.7, "y": 2.9, "width": 1.4, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-table-non-first-rowspan.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 3
    assert len(table.columns) == 3
    assert table.cell(1, 1).text == "APAC + EMEA"
    assert table.cell(1, 2).text == "$2.4M"
    assert table.cell(2, 2).text == "$1.8M"
    assert 'rowSpan="2"' in table.cell(1, 1)._tc.xml


def test_build_editable_rebuild_styles_second_header_row_in_multilevel_tables(tmp_path):
    blocks = [
        {"text": "Business review", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.2, "height": 0.6, "font_size": 26},
        {"text": "Commercial", "slide_index": 0, "x": 1.0, "y": 1.5, "width": 4.9, "height": 0.45, "font_size": 18},
        {"text": "Operations", "slide_index": 0, "x": 5.7, "y": 1.5, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Revenue", "slide_index": 0, "x": 3.5, "y": 2.0, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "Utilization", "slide_index": 0, "x": 5.7, "y": 2.0, "width": 1.4, "height": 0.45, "font_size": 18},
        {"text": "APAC", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "$2.4M", "slide_index": 0, "x": 3.5, "y": 2.7, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "81%", "slide_index": 0, "x": 5.7, "y": 2.7, "width": 1.4, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-multilevel-header.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    top_header_run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    second_header_run = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
    body_run = table.cell(2, 0).text_frame.paragraphs[0].runs[0]
    second_header_fill = table.cell(1, 0).fill.fore_color.rgb
    body_fill = table.cell(2, 0).fill.fore_color.rgb

    assert top_header_run.font.bold is True
    assert second_header_run.font.bold is True
    assert body_run.font.bold in (None, False)
    assert second_header_fill == RGBColor(0xE9, 0xEE, 0xF7)
    assert body_fill != second_header_fill
    assert 'gridSpan="2"' in table.cell(0, 0)._tc.xml


def test_build_editable_rebuild_handles_constrained_cross_merges(tmp_path):
    blocks = [
        {"text": "Business review", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.4, "height": 0.6, "font_size": 26},
        {"text": "Commercial", "slide_index": 0, "x": 1.0, "y": 1.5, "width": 4.8, "height": 0.45, "font_size": 18},
        {"text": "Operations", "slide_index": 0, "x": 5.6, "y": 1.5, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Category", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 3.4, "y": 2.0, "width": 1.7, "height": 0.45, "font_size": 18},
        {"text": "Utilization", "slide_index": 0, "x": 5.6, "y": 2.0, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.0, "height": 1.0, "font_size": 18},
        {"text": "APAC + EMEA", "slide_index": 0, "x": 3.4, "y": 2.7, "width": 1.7, "height": 1.0, "font_size": 18},
        {"text": "81%", "slide_index": 0, "x": 5.6, "y": 2.7, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "79%", "slide_index": 0, "x": 5.6, "y": 3.3, "width": 1.6, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-constrained-cross-merge.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 4
    assert len(table.columns) == 3
    assert 'gridSpan="2"' in table.cell(0, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 1)._tc.xml
    assert table.cell(2, 2).text == "81%"
    assert table.cell(3, 2).text == "79%"


def test_build_editable_rebuild_handles_multiple_non_overlapping_cross_merges(tmp_path):
    blocks = [
        {"text": "Business review", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 6.6, "height": 0.6, "font_size": 26},
        {"text": "Commercial", "slide_index": 0, "x": 1.0, "y": 1.5, "width": 4.9, "height": 0.45, "font_size": 18},
        {"text": "Operations", "slide_index": 0, "x": 5.8, "y": 1.5, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Category", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 3.4, "y": 2.0, "width": 1.8, "height": 0.45, "font_size": 18},
        {"text": "Utilization", "slide_index": 0, "x": 5.8, "y": 2.0, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.0, "height": 1.0, "font_size": 18},
        {"text": "APAC + EMEA", "slide_index": 0, "x": 3.4, "y": 2.7, "width": 1.8, "height": 1.0, "font_size": 18},
        {"text": "81%", "slide_index": 0, "x": 5.8, "y": 2.7, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "79%", "slide_index": 0, "x": 5.8, "y": 3.3, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Consumer", "slide_index": 0, "x": 1.0, "y": 4.0, "width": 2.0, "height": 1.0, "font_size": 18},
        {"text": "Americas + LATAM", "slide_index": 0, "x": 3.4, "y": 4.0, "width": 1.8, "height": 1.0, "font_size": 18},
        {"text": "74%", "slide_index": 0, "x": 5.8, "y": 4.0, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "71%", "slide_index": 0, "x": 5.8, "y": 4.6, "width": 1.6, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-multi-cross-merge.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 6
    assert len(table.columns) == 3
    assert 'gridSpan="2"' in table.cell(0, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 1)._tc.xml
    assert 'rowSpan="2"' in table.cell(4, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(4, 1)._tc.xml
    assert table.cell(5, 2).text == "71%"


def test_build_editable_rebuild_handles_freer_cross_merge_combinations_without_conflicts(tmp_path):
    blocks = [
        {"text": "Portfolio review", "slide_index": 0, "x": 1.0, "y": 0.8, "width": 7.4, "height": 0.6, "font_size": 26},
        {"text": "Commercial", "slide_index": 0, "x": 1.0, "y": 1.5, "width": 4.8, "height": 0.45, "font_size": 18},
        {"text": "Operations", "slide_index": 0, "x": 5.8, "y": 1.5, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Customer", "slide_index": 0, "x": 7.7, "y": 1.5, "width": 1.5, "height": 0.45, "font_size": 18},
        {"text": "Category", "slide_index": 0, "x": 1.0, "y": 2.0, "width": 2.0, "height": 0.45, "font_size": 18},
        {"text": "Region", "slide_index": 0, "x": 3.4, "y": 2.0, "width": 1.8, "height": 0.45, "font_size": 18},
        {"text": "Utilization", "slide_index": 0, "x": 5.8, "y": 2.0, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "NPS", "slide_index": 0, "x": 7.7, "y": 2.0, "width": 1.5, "height": 0.45, "font_size": 18},
        {"text": "Enterprise", "slide_index": 0, "x": 1.0, "y": 2.7, "width": 2.0, "height": 1.0, "font_size": 18},
        {"text": "APAC + EMEA", "slide_index": 0, "x": 3.4, "y": 2.7, "width": 1.8, "height": 1.0, "font_size": 18},
        {"text": "81%", "slide_index": 0, "x": 5.8, "y": 2.7, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "57", "slide_index": 0, "x": 7.7, "y": 2.7, "width": 1.5, "height": 1.0, "font_size": 18},
        {"text": "79%", "slide_index": 0, "x": 5.8, "y": 3.3, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "Consumer", "slide_index": 0, "x": 1.0, "y": 4.0, "width": 2.0, "height": 1.0, "font_size": 18},
        {"text": "Americas + LATAM", "slide_index": 0, "x": 3.4, "y": 4.0, "width": 1.8, "height": 1.0, "font_size": 18},
        {"text": "74%", "slide_index": 0, "x": 5.8, "y": 4.0, "width": 1.6, "height": 0.45, "font_size": 18},
        {"text": "49", "slide_index": 0, "x": 7.7, "y": 4.0, "width": 1.5, "height": 1.0, "font_size": 18},
        {"text": "71%", "slide_index": 0, "x": 5.8, "y": 4.6, "width": 1.6, "height": 0.45, "font_size": 18},
    ]
    output_path = tmp_path / "editable-rebuild-freer-cross-merge.pptx"

    build_editable_rebuild(blocks, output_path)

    presentation = Presentation(output_path)
    table = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.TABLE][0].table

    assert len(table.rows) == 6
    assert len(table.columns) == 4
    assert 'gridSpan="2"' in table.cell(0, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 1)._tc.xml
    assert 'rowSpan="2"' in table.cell(2, 3)._tc.xml
    assert 'rowSpan="2"' in table.cell(4, 0)._tc.xml
    assert 'rowSpan="2"' in table.cell(4, 1)._tc.xml
    assert 'rowSpan="2"' in table.cell(4, 3)._tc.xml
    assert table.cell(5, 2).text == "71%"
