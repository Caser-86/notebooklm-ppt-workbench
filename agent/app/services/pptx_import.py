import base64
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation


PLACEHOLDER_PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wn7Zp4AAAAASUVORK5CYII="
)


@dataclass
class ImportedSlideBundle:
    slide_index: int
    preview_image_path: str
    text_dump: str
    structure_json_path: str


@dataclass
class ImportedPresentationBundle:
    source_type: str
    page_count: int
    slides: list[ImportedSlideBundle]


def classify_pptx_source(pptx_path: Path) -> str:
    presentation = Presentation(pptx_path)
    texts = " ".join(_iter_slide_texts(presentation)).lower()
    if "editable rebuild" in texts or "display clone" in texts:
        return "internal_generated"
    if "notebooklm" in texts:
        return "notebooklm_export"
    return "generic_pptx"


def extract_pptx_assets(project_id: int, pptx_path: Path, import_dir: Path) -> ImportedPresentationBundle:
    del project_id  # project_id is reserved for future slide-specific asset strategies.
    presentation = Presentation(pptx_path)
    source_type = classify_pptx_source(pptx_path)
    slides: list[ImportedSlideBundle] = []

    for index, slide in enumerate(presentation.slides, start=1):
        text_dump = "\n".join(
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text and shape.text.strip()
        )
        preview_path = import_dir / f"slide-{index}.png"
        preview_path.write_bytes(PLACEHOLDER_PNG_BYTES)
        slides.append(
            ImportedSlideBundle(
                slide_index=index,
                preview_image_path=str(preview_path),
                text_dump=text_dump,
                structure_json_path="",
            )
        )

    return ImportedPresentationBundle(source_type=source_type, page_count=len(slides), slides=slides)


def _iter_slide_texts(presentation: Presentation) -> list[str]:
    texts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
    return texts
