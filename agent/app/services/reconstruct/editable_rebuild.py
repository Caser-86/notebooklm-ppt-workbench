from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.services.reconstruct.layout_analysis import analyze_rebuild_layout


def resolve_render_geometry(block: dict) -> tuple[float, float, float, float]:
    left = block["x"]
    top = block["y"]
    width = block["width"]
    height = block["height"]

    if block.get("text_role") in {"body", "list_item", "caption"}:
        column_index = block.get("column_index")
        if column_index == 0:
            width = min(width, 5.1)
        elif column_index == 1:
            left = max(left, 6.8)
            width = min(width, 5.1)

    return left, top, width, height


def set_table_cell_borders(cell, color: str = "A8B3C7", width: str = "12700") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for border_name in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        existing = tc_pr.find(qn(border_name))
        if existing is not None:
            tc_pr.remove(existing)

        border = OxmlElement(border_name)
        border.set("w", width)
        solid_fill = OxmlElement("a:solidFill")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color)
        solid_fill.append(srgb)
        border.append(solid_fill)
        preset_dash = OxmlElement("a:prstDash")
        preset_dash.set("val", "solid")
        border.append(preset_dash)
        round_join = OxmlElement("a:round")
        border.append(round_join)
        head_end = OxmlElement("a:headEnd")
        head_end.set("type", "none")
        head_end.set("w", "med")
        head_end.set("len", "med")
        border.append(head_end)
        tail_end = OxmlElement("a:tailEnd")
        tail_end.set("type", "none")
        tail_end.set("w", "med")
        tail_end.set("len", "med")
        border.append(tail_end)
        tc_pr.append(border)


def style_table_cell(cell, is_header: bool) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0xE9, 0xEE, 0xF7) if is_header else RGBColor(0xFF, 0xFF, 0xFF)
    set_table_cell_borders(cell)


def build_editable_rebuild(raw_blocks: list[dict], output_path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    if any(str(block.get("content_type", "")).startswith("imported_") for block in raw_blocks):
        slide_groups = _group_imported_blocks_by_slide(raw_blocks)
    else:
        slide_groups = analyze_rebuild_layout(raw_blocks)

    for slide_blocks in slide_groups:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for block in slide_blocks:
            left, top, width, height = resolve_render_geometry(block)
            if block.get("content_type") in {"table", "imported_table"}:
                table_shape = slide.shapes.add_table(
                    block["rows"],
                    block["cols"],
                    Inches(left),
                    Inches(top),
                    Inches(width),
                    Inches(height),
                )
                table = table_shape.table
                for column_index, column_width in enumerate(block.get("column_widths", [])):
                    table.columns[column_index].width = Inches(column_width)
                header_rows = int(block.get("header_rows", 1))
                for row_index, row_cells in enumerate(block["cells"]):
                    for column_index, cell_data in enumerate(row_cells):
                        if cell_data.get("merged"):
                            continue
                        cell = table.cell(row_index, column_index)
                        colspan = int(cell_data.get("colspan", 1))
                        rowspan = int(cell_data.get("rowspan", 1))
                        if colspan > 1 or rowspan > 1:
                            cell.merge(table.cell(row_index + rowspan - 1, column_index + colspan - 1))
                        style_table_cell(cell, is_header=row_index < header_rows)
                        text_frame = cell.text_frame
                        text_frame.clear()
                        paragraph = text_frame.paragraphs[0]
                        run = paragraph.add_run()
                        run.text = cell_data["text"]
                        run.font.size = Pt(cell_data.get("font_size", 18))
                        run.font.bold = row_index < header_rows
                continue

            if block.get("content_type") in {"icon_card", "imported_icon_card"}:
                slide.shapes.add_picture(
                    str(Path(block["icon_path"])),
                    left=Inches(block["icon_x"]),
                    top=Inches(block["icon_y"]),
                    width=Inches(block["icon_width"]),
                    height=Inches(block["icon_height"]),
                )

                title_box = slide.shapes.add_textbox(
                    left=Inches(block["title_x"]),
                    top=Inches(block["title_y"]),
                    width=Inches(block["title_width"]),
                    height=Inches(block["title_height"]),
                )
                title_frame = title_box.text_frame
                title_frame.clear()
                title_run = title_frame.paragraphs[0].add_run()
                title_run.text = block["title_text"]
                title_run.font.size = Pt(block["title_font_size"])
                title_run.font.bold = True

                body_box = slide.shapes.add_textbox(
                    left=Inches(block["body_x"]),
                    top=Inches(block["body_y"]),
                    width=Inches(block["body_width"]),
                    height=Inches(block["body_height"]),
                )
                body_frame = body_box.text_frame
                body_frame.clear()
                body_run = body_frame.paragraphs[0].add_run()
                body_run.text = block["body_text"]
                body_run.font.size = Pt(block["body_font_size"])
                continue

            if block.get("content_type") in {"image", "imported_image"}:
                image_path = Path(block["image_path"])
                slide.shapes.add_picture(
                    str(image_path),
                    left=Inches(left),
                    top=Inches(top),
                    width=Inches(width),
                    height=Inches(height),
                )
                continue

            textbox = slide.shapes.add_textbox(
                left=Inches(left),
                top=Inches(top),
                width=Inches(width),
                height=Inches(height),
            )
            text_frame = textbox.text_frame
            text_frame.clear()

            paragraphs = block.get("paragraphs", [block])
            for index, paragraph_block in enumerate(paragraphs):
                paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
                paragraph.level = 1 if paragraph_block.get("text_role") == "list_item" else 0
                run = paragraph.add_run()
                run.text = paragraph_block["text"]
                run.font.size = Pt(paragraph_block.get("font_size", 18))
                run.font.bold = (
                    paragraph_block.get("bold")
                    if paragraph_block.get("bold") is not None
                    else paragraph_block.get("text_role") == "title"
                )
                run.font.italic = (
                    paragraph_block.get("italic")
                    if paragraph_block.get("italic") is not None
                    else paragraph_block.get("text_role") == "caption"
                )

    presentation.save(output_path)
    return output_path


def _group_imported_blocks_by_slide(raw_blocks: list[dict]) -> list[list[dict]]:
    grouped: dict[int, list[dict]] = {}
    for block in raw_blocks:
        slide_index = int(block.get("slide_index", 0))
        grouped.setdefault(slide_index, []).append(dict(block))

    return [
        sorted(grouped[slide_index], key=lambda item: (item.get("y", 0), item.get("x", 0)))
        for slide_index in sorted(grouped)
    ]
