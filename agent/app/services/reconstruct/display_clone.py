from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def build_display_clone(slide_paths: list[Path], output_path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    blank_layout = presentation.slide_layouts[6]
    for slide_path in slide_paths:
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(str(slide_path), 0, 0, width=presentation.slide_width, height=presentation.slide_height)

    if presentation.slides:
        first = presentation.slides._sldIdLst[0]
        presentation.slides._sldIdLst.remove(first)

    presentation.save(output_path)
    return output_path
